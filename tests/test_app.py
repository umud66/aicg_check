from io import BytesIO

from docx import Document

from app import app


def make_docx() -> BytesIO:
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph(
        "首先，应进一步完善基层治理机制。其次，应加强部门之间的协同联动，推动治理资源有效整合。"
        "此外，还应不断优化公共服务体系，从而进一步提升基层治理效能。"
    )
    document.add_paragraph(
        "根据2024年公开数据，样本中共有53所学校，其中31所为小学，22所为初中[1]。"
        "这一事实用于说明样本结构，而不是泛化提出政策建议。"
    )
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def test_healthz():
    client = app.test_client()
    response = client.get("/healthz")
    assert response.status_code == 200
    assert response.json == {"status": "ok"}


def test_docx_analysis():
    client = app.test_client()
    response = client.post(
        "/api/analyze",
        data={
            "file": (make_docx(), "sample.docx"),
            "medium_threshold": "40",
            "high_threshold": "65",
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    payload = response.json
    assert payload["summary"]["total_chars"] > 0
    assert len(payload["results"]) == 2
    assert payload["results"][0]["chapter"] == "第一章 绪论"


def test_rejects_bad_extension():
    client = app.test_client()
    response = client.post(
        "/api/analyze",
        data={"file": (BytesIO(b"hello"), "sample.txt")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
